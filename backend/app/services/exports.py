"""Asynchronous, ACL-protected transcript exports."""

from __future__ import annotations

import io
import json
import os
import re
import tempfile
from pathlib import Path
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ApiError
from app.models.enums import ExportFormat, JobStatus, MeetingPermission
from app.models.export import Export
from app.models.media import Job
from app.models.meeting import Meeting
from app.models.summary import Summary, SummaryVersion
from app.models.transcript import Speaker, TranscriptVersion
from app.models.translation import Translation, TranslationSegment
from app.models.user import User
from app.providers.storage import MediaStorage, StorageError
from app.schemas.export import ExportCreateRequest, ExportResponse
from app.services.meetings import MeetingService
from app.services.transcript_versions import TranscriptVersionReader

_CONTENT_TYPES = {
    ExportFormat.TXT: "text/plain; charset=utf-8",
    ExportFormat.SRT: "application/x-subrip; charset=utf-8",
    ExportFormat.VTT: "text/vtt; charset=utf-8",
    ExportFormat.JSON: "application/json; charset=utf-8",
    ExportFormat.PDF: "application/pdf",
    ExportFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExportService:
    def __init__(self, session: AsyncSession, storage: MediaStorage) -> None:
        self._session, self._storage = session, storage
        self._meetings = MeetingService(session)

    async def request(
        self, meeting_id: UUID, user: User, payload: ExportCreateRequest
    ) -> ExportResponse:
        meeting, _ = await self._meetings.require_access(meeting_id, user, MeetingPermission.READ)
        version = await self._current_version(meeting)
        export = Export(
            meeting_id=meeting.id, transcript_version_id=version.id, format=payload.format
        )
        self._session.add(export)
        await self._session.flush()
        job = Job(
            meeting_id=meeting.id,
            job_type="export",
            status=JobStatus.QUEUED,
            progress=0,
            payload={"export_id": str(export.id), "options": payload.options.model_dump()},
        )
        self._session.add(job)
        await self._session.commit()
        await self._session.refresh(export)
        await self._session.refresh(job)
        return self._response(export, job)

    async def get(self, export_id: UUID, user: User) -> ExportResponse:
        export = await self._export_for_read(export_id, user)
        return self._response(export, await self._job(export.id))

    async def download(self, export_id: UUID, user: User) -> tuple[bytes, Export]:
        export = await self._export_for_read(export_id, user)
        job = await self._job(export.id)
        if job is None or job.status != JobStatus.COMPLETED or not export.storage_key:
            raise ApiError("EXPORT_NOT_READY", "The export is not ready to download.", 409)
        try:
            return await self._storage.get_bytes(export.storage_key), export
        except (OSError, StorageError) as exc:
            raise ApiError("EXPORT_STORAGE_ERROR", "The export file is unavailable.", 503) from exc

    async def presigned_download_url(
        self, export_id: UUID, user: User, expires_in: int
    ) -> str | None:
        export = await self._export_for_read(export_id, user)
        job = await self._job(export.id)
        if job is None or job.status != JobStatus.COMPLETED or not export.storage_key:
            raise ApiError("EXPORT_NOT_READY", "The export is not ready to download.", 409)
        disposition = f'attachment; filename="{export.filename or "meeting-export"}"'
        try:
            return await self._storage.presigned_get_url(
                export.storage_key,
                expires_in,
                content_type=export.content_type,
                content_disposition=disposition,
            )
        except StorageError as exc:
            raise ApiError(
                "EXPORT_STORAGE_ERROR", "The export URL could not be created.", 503
            ) from exc

    async def purge_meeting_storage(self, meeting_id: UUID) -> None:
        keys = list(
            await self._session.scalars(
                select(Export.storage_key).where(
                    Export.meeting_id == meeting_id, Export.storage_key.is_not(None)
                )
            )
        )
        try:
            for key in keys:
                await self._storage.delete(key)
        except StorageError as exc:
            raise ApiError(
                "EXPORT_STORAGE_ERROR", "Meeting exports could not be deleted.", 502
            ) from exc

    async def _current_version(self, meeting: Meeting) -> TranscriptVersion:
        if meeting.current_transcript_version_id is None:
            raise ApiError("TRANSCRIPT_NOT_FOUND", "A transcript is required for export.", 422)
        version = await self._session.get(TranscriptVersion, meeting.current_transcript_version_id)
        if version is None:
            raise ApiError("TRANSCRIPT_NOT_FOUND", "A transcript is required for export.", 422)
        return version

    async def _export_for_read(self, export_id: UUID, user: User) -> Export:
        export = await self._session.get(Export, export_id)
        if export is None:
            raise ApiError("EXPORT_NOT_FOUND", "Export was not found.", 404)
        await self._meetings.require_access(export.meeting_id, user, MeetingPermission.READ)
        return export

    async def _job(self, export_id: UUID) -> Job | None:
        jobs = list(await self._session.scalars(select(Job).where(Job.job_type == "export")))
        return next(
            (job for job in jobs if (job.payload or {}).get("export_id") == str(export_id)), None
        )

    @staticmethod
    def _response(export: Export, job: Job | None) -> ExportResponse:
        return ExportResponse(
            id=export.id,
            meeting_id=export.meeting_id,
            transcript_version_id=export.transcript_version_id,
            format=export.format,
            job_id=job.id if job else None,
            status=job.status if job else JobStatus.FAILED,
            filename=export.filename,
            content_type=export.content_type,
            created_at=export.created_at,
        )


class ExportJobRunner:
    """Runs export jobs using the same durable ``jobs`` table as AI jobs."""

    def __init__(self, storage: MediaStorage) -> None:
        self._storage = storage

    async def run(self, session: AsyncSession, job: Job) -> None:
        payload = job.payload or {}
        export = await session.get(Export, UUID(str(payload["export_id"])))
        if export is None:
            raise ValueError("Export request no longer exists.")
        meeting = await session.get(Meeting, export.meeting_id)
        if meeting is None:
            raise ValueError("Meeting no longer exists.")
        version = await session.get(TranscriptVersion, export.transcript_version_id)
        if version is None:
            raise ValueError("Transcript version no longer exists.")
        try:
            segments = await TranscriptVersionReader(session).read(version)
        except ApiError as exc:
            # Missing history is terminal input failure, not a transient provider failure.
            raise ValueError(str(exc)) from exc
        if not segments:
            raise ValueError("Transcript input is empty or invalid.")
        options = {
            "include_timestamps": True,
            "include_speakers": True,
            "include_translation": False,
            "include_summary": False,
            **dict(payload.get("options", {})),
        }
        speakers = {
            item.id: item.name
            for item in await session.scalars(
                select(Speaker).where(Speaker.meeting_id == meeting.id)
            )
        }
        translation = (
            await _translation_by_segment(session, meeting.id, export.transcript_version_id)
            if options["include_translation"]
            else {}
        )
        summary = (
            await _latest_summary(session, meeting.id, export.transcript_version_id)
            if options["include_summary"]
            else None
        )
        document = _document(meeting, export, segments, speakers, translation, summary, options)
        content = _render(export.format, document)
        filename = _filename(meeting.title, export.format)
        key = f"exports/meetings/{meeting.id}/{export.id}.{export.format.value}"
        # Export binaries can be large. Keep the object-storage boundary
        # streaming-friendly and avoid assuming an API/worker shared volume.
        descriptor, temporary_name = tempfile.mkstemp(
            prefix="bocbang-export-", suffix=f".{export.format.value}"
        )
        temporary = Path(temporary_name)
        try:
            with os.fdopen(descriptor, "wb") as output:
                output.write(content)
            await self._storage.put_file(temporary, key)
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
        export.storage_key, export.filename, export.content_type = (
            key,
            filename,
            _CONTENT_TYPES[export.format],
        )


async def _translation_by_segment(
    session: AsyncSession, meeting_id: UUID, version_id: UUID
) -> dict[UUID, str]:
    translation = await session.scalar(
        select(Translation)
        .where(
            Translation.meeting_id == meeting_id, Translation.transcript_version_id == version_id
        )
        .order_by(Translation.created_at.desc())
        .limit(1)
    )
    if translation is None:
        return {}
    rows = await session.scalars(
        select(TranslationSegment).where(TranslationSegment.translation_id == translation.id)
    )
    return {row.logical_segment_id: row.text for row in rows}


async def _latest_summary(
    session: AsyncSession, meeting_id: UUID, version_id: UUID
) -> dict[str, object] | None:
    summary = await session.scalar(
        select(Summary)
        .where(Summary.meeting_id == meeting_id, Summary.transcript_version_id == version_id)
        .order_by(Summary.created_at.desc())
        .limit(1)
    )
    if summary is None:
        return None
    row = await session.scalar(
        select(SummaryVersion)
        .where(SummaryVersion.summary_id == summary.id)
        .order_by(SummaryVersion.version_no.desc())
        .limit(1)
    )
    return row.content if row else None


def _document(
    meeting: Meeting,
    export: Export,
    segments: list[dict[str, object]],
    speakers: dict[UUID, str],
    translation: dict[UUID, str],
    summary: dict[str, object] | None,
    options: dict[str, object],
) -> dict[str, object]:
    data = {
        "meeting": {
            "title": meeting.title,
            "description": meeting.description,
            "language_mode": meeting.language_mode,
            "created_at": meeting.created_at.isoformat(),
            "transcript_version_id": str(export.transcript_version_id),
        },
        "options": options,
        "segments": [],
    }
    for item in segments:
        row: dict[str, object] = {
            "id": str(item["logical_segment_id"]),
            "start_ms": item["start_ms"],
            "end_ms": item["end_ms"],
            "text": item["text"],
        }
        if options["include_speakers"]:
            row["speaker"] = speakers.get(item["speaker_id"]) if item["speaker_id"] else None
        if options["include_translation"] and item["logical_segment_id"] in translation:
            row["translation"] = translation[item["logical_segment_id"]]
        data["segments"].append(row)
    if options["include_summary"] and summary is not None:
        data["summary"] = summary
    return data


def _render(format: ExportFormat, data: dict[str, object]) -> bytes:
    if format == ExportFormat.JSON:
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
    if format in (ExportFormat.SRT, ExportFormat.VTT):
        return _subtitles(data, format).encode("utf-8")
    text = _plain_text(data)
    if format == ExportFormat.TXT:
        return text.encode("utf-8")
    if format == ExportFormat.DOCX:
        from docx import Document

        document = Document()
        document.add_heading(str(data["meeting"]["title"]), 0)
        for line in text.splitlines():
            document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    # DejaVu is shipped by many Linux images; failure is a stable failed job rather than mojibake.
    font = next(
        (
            path
            for path in (
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                "C:/Windows/Fonts/arial.ttf",
            )
            if os.path.exists(path)
        ),
        None,
    )
    if font is None:
        raise ValueError("A Unicode PDF font is not configured.")
    pdfmetrics.registerFont(TTFont("ExportUnicode", font))
    buffer = io.BytesIO()
    canvas = Canvas(buffer, pagesize=A4)
    canvas.setFont("ExportUnicode", 11)
    y = 800
    for line in text.splitlines():
        for part in _wrap(line, 95):
            canvas.drawString(42, y, part)
            y -= 15
            if y < 42:
                canvas.showPage()
                canvas.setFont("ExportUnicode", 11)
                y = 800
    canvas.save()
    return buffer.getvalue()


def _plain_text(data: dict[str, object]) -> str:
    meeting = data["meeting"]
    lines = [str(meeting["title"]), "", f"Language: {meeting['language_mode']}"]
    if meeting.get("description"):
        lines.append(str(meeting["description"]))
    for item in data["segments"]:
        prefix = (
            f"[{_timestamp(item['start_ms'])} - {_timestamp(item['end_ms'])}] "
            if data["options"]["include_timestamps"]
            else ""
        )
        if item.get("speaker"):
            prefix += f"{item['speaker']}: "
        lines.append(prefix + str(item["text"]))
        if item.get("translation"):
            lines.append(f"  Translation: {item['translation']}")
    if data.get("summary"):
        lines.extend(["", "Summary", json.dumps(data["summary"], ensure_ascii=False)])
    return "\n".join(lines) + "\n"


def _subtitles(data: dict[str, object], format: ExportFormat) -> str:
    lines = ["WEBVTT", ""] if format == ExportFormat.VTT else []
    for index, item in enumerate(data["segments"], start=1):
        if format == ExportFormat.SRT:
            lines.append(str(index))
        separator = "." if format == ExportFormat.VTT else ","
        lines.append(
            f"{_timestamp(item['start_ms'], separator)} --> {_timestamp(item['end_ms'], separator)}"
        )
        label = f"{item['speaker']}: " if item.get("speaker") else ""
        lines.append(label + str(item["text"]))
        lines.append("")
    return "\n".join(lines)


def _timestamp(value: object, separator: str = ".") -> str:
    ms = int(value)
    hours, ms = divmod(ms, 3_600_000)
    minutes, ms = divmod(ms, 60_000)
    seconds, ms = divmod(ms, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02}{separator}{ms:03}"


def _filename(title: str, format: ExportFormat) -> str:
    stem = re.sub(r"[^\w.-]+", "-", title, flags=re.UNICODE).strip(".-") or "meeting"
    return f"{stem}.{format.value}"


def _wrap(value: str, width: int) -> list[str]:
    return [value[index : index + width] for index in range(0, len(value), width)] or [""]
